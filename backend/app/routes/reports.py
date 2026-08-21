import csv
import io
from datetime import date

from flask import Blueprint, Response, jsonify, request

from ..extensions import db
from ..models import Lease, MaintenanceRequest, Payment, Property, Tenant
from ..utils import parse_date, staff_required

bp = Blueprint("reports", __name__, url_prefix="/api/reports")

MONEY_COLUMNS = {"Rent", "Billed", "Paid", "Balance", "Deposit", "Cost", "Amount", "Rent roll"}

REPORT_TYPES = (
    {"value": "rent", "label": "Rent Collection Report", "description": "Monthly rent billed, collected and outstanding"},
    {"value": "payments", "label": "Payment Report", "description": "Every transaction with method and status"},
    {"value": "properties", "label": "Property Report", "description": "Occupancy and availability across the portfolio"},
    {"value": "tenants", "label": "Tenant Report", "description": "Tenant directory with lease status"},
    {"value": "leases", "label": "Lease Report", "description": "Agreements, terms and rent roll"},
    {"value": "maintenance", "label": "Maintenance Report", "description": "Request volume, ageing and cost"},
)


def _payments_in_range(start, end, status=None, prop=None, tenant=None):
    query = Payment.query.join(Lease)
    if start:
        query = query.filter(Payment.due_date >= start)
    if end:
        query = query.filter(Payment.due_date <= end)
    if status and status != "all":
        query = query.filter(Payment.status == status)
    if prop:
        query = query.filter(Lease.property_id == prop)
    if tenant:
        query = query.filter(Lease.tenant_id == tenant)
    return query.order_by(Payment.due_date).all()


def build(report, filters):
    """Return (columns, rows, totals) for the requested report."""
    start = filters.get("start")
    end = filters.get("end")
    prop = filters.get("property_id")
    tenant = filters.get("tenant_id")
    status = filters.get("status")

    if report == "properties":
        query = Property.query
        if status and status != "all":
            query = query.filter(Property.status == status)
        if prop:
            query = query.filter(Property.id == prop)
        items = query.order_by(Property.name).all()
        columns = ["Property ID", "Name", "Address", "Type", "Furnishing", "Rent", "Status", "Current tenant"]
        rows = [
            [
                p.property_code or f"PR-{p.id:04d}",
                p.name,
                ", ".join(filter(None, [p.address, p.city])),
                p.property_type,
                p.furnishing or "-",
                float(p.rent_amount or 0),
                p.status,
                p.to_dict()["current_tenant"] or "-",
            ]
            for p in items
        ]
        totals = {
            "Properties": len(items),
            "Occupied": sum(1 for p in items if p.status == "occupied"),
            "Available": sum(1 for p in items if p.status == "available"),
            "Occupancy %": round(
                sum(1 for p in items if p.status == "occupied") / len(items) * 100, 1
            ) if items else 0,
            "Rent roll": sum(float(p.rent_amount or 0) for p in items if p.status == "occupied"),
        }

    elif report == "tenants":
        query = Tenant.query
        if tenant:
            query = query.filter(Tenant.id == tenant)
        items = query.order_by(Tenant.full_name).all()
        columns = ["Name", "Email", "Phone", "ID proof", "Property", "Unit", "Lease status"]
        rows = []
        for t in items:
            info = t.to_dict()
            rows.append([
                t.full_name,
                t.email,
                t.phone or "-",
                t.id_proof_type or "-",
                info["current_property"] or "-",
                t.unit_room or "-",
                info["lease_status"],
            ])
        totals = {
            "Tenants": len(items),
            "With active lease": sum(1 for t in items if t.active_lease),
            "Unassigned": sum(1 for t in items if not t.active_lease),
            "With portal login": sum(1 for t in items if t.user_id),
        }

    elif report == "leases":
        query = Lease.query
        if prop:
            query = query.filter(Lease.property_id == prop)
        if tenant:
            query = query.filter(Lease.tenant_id == tenant)
        if start:
            query = query.filter(Lease.end_date >= start)
        if end:
            query = query.filter(Lease.start_date <= end)
        items = query.order_by(Lease.start_date.desc()).all()
        columns = ["Lease ID", "Property", "Tenant", "Start", "End", "Rent", "Deposit", "Status"]
        rows = [
            [
                l.lease_code or f"LS-{l.id:04d}",
                l.property.name if l.property else "-",
                l.tenant.full_name if l.tenant else "-",
                l.start_date.isoformat(),
                l.end_date.isoformat(),
                float(l.rent_amount or 0),
                float(l.deposit_amount or 0),
                l.lease_state,
            ]
            for l in items
        ]
        totals = {
            "Leases": len(items),
            "Active": sum(1 for l in items if l.status == "active"),
            "Expiring soon": sum(1 for l in items if l.lease_state == "expiring_soon"),
            "Rent roll": sum(float(l.rent_amount or 0) for l in items if l.status == "active"),
            "Deposit": sum(float(l.deposit_amount or 0) for l in items if l.status == "active"),
        }

    elif report == "maintenance":
        query = MaintenanceRequest.query
        if prop:
            query = query.filter(MaintenanceRequest.property_id == prop)
        if tenant:
            query = query.filter(MaintenanceRequest.tenant_id == tenant)
        if status and status != "all":
            query = query.filter(MaintenanceRequest.status == status)
        if start:
            query = query.filter(MaintenanceRequest.created_at >= start)
        if end:
            query = query.filter(MaintenanceRequest.created_at <= end)
        items = query.order_by(MaintenanceRequest.created_at.desc()).all()
        columns = ["Ticket", "Property", "Tenant", "Title", "Category", "Priority", "Status", "Days", "Cost", "Raised"]
        rows = [
            [
                m.ticket_code or f"MR-{m.id:04d}",
                m.property.name if m.property else "-",
                m.tenant.full_name if m.tenant else "-",
                m.title,
                m.category,
                m.priority,
                m.status,
                m.age_days,
                float(m.cost or 0),
                m.created_at.strftime("%Y-%m-%d") if m.created_at else "-",
            ]
            for m in items
        ]
        resolved = [m for m in items if m.status in ("completed", "closed")]
        totals = {
            "Requests": len(items),
            "Open": sum(1 for m in items if m.is_open),
            "Completed": len(resolved),
            "Avg days to resolve": round(
                sum(m.age_days for m in resolved) / len(resolved), 1
            ) if resolved else 0,
            "Cost": sum(float(m.cost or 0) for m in items),
        }

    elif report == "payments":
        items = _payments_in_range(start, end, status, prop, tenant)
        columns = ["Payment ID", "Tenant", "Property", "Due date", "Paid date", "Amount", "Method", "Reference", "Status"]
        rows = [
            [
                p.payment_code or f"PY-{p.id:05d}",
                p.lease.tenant.full_name if p.lease and p.lease.tenant else "-",
                p.lease.property.name if p.lease and p.lease.property else "-",
                p.due_date.isoformat() if p.due_date else "-",
                p.paid_date.isoformat() if p.paid_date else "-",
                float(p.amount_paid or 0),
                (p.method or "-").replace("_", " "),
                p.reference or "-",
                p.status,
            ]
            for p in items
        ]
        totals = {
            "Transactions": len(items),
            "Amount": sum(float(p.amount_paid or 0) for p in items),
            "Paid in full": sum(1 for p in items if p.status == "paid"),
            "Part paid": sum(1 for p in items if p.status == "partial"),
        }

    else:  # rent collection
        items = _payments_in_range(start, end, status, prop, tenant)
        columns = ["Period", "Property", "Tenant", "Due date", "Billed", "Paid", "Balance", "Status"]
        rows = [
            [
                p.period or "-",
                p.lease.property.name if p.lease and p.lease.property else "-",
                p.lease.tenant.full_name if p.lease and p.lease.tenant else "-",
                p.due_date.isoformat() if p.due_date else "-",
                float(p.amount or 0),
                float(p.amount_paid or 0),
                float(p.amount or 0) - float(p.amount_paid or 0),
                p.status,
            ]
            for p in items
        ]
        billed = sum(float(p.amount or 0) for p in items)
        collected = sum(float(p.amount_paid or 0) for p in items)
        totals = {
            "Invoices": len(items),
            "Billed": billed,
            "Paid": collected,
            "Balance": billed - collected,
            "Collection %": round(collected / billed * 100, 1) if billed else 0,
        }

    return columns, rows, totals


def read_filters(args):
    return {
        "start": parse_date(args.get("start_date"), "start_date"),
        "end": parse_date(args.get("end_date"), "end_date"),
        "property_id": int(args["property_id"]) if args.get("property_id") else None,
        "tenant_id": int(args["tenant_id"]) if args.get("tenant_id") else None,
        "status": args.get("status"),
    }


@bp.get("/types")
@staff_required
def report_types():
    return jsonify(list(REPORT_TYPES))


@bp.get("/dashboard")
@staff_required
def reports_dashboard():
    """Reports Dashboard — headline numbers plus what each report covers."""
    today = date.today()
    properties = Property.query.all()
    leases = Lease.query.all()
    payments = Payment.query.filter(Payment.due_date <= today).all()
    requests = MaintenanceRequest.query.all()

    billed = sum(float(p.amount or 0) for p in payments)
    collected = sum(float(p.amount_paid or 0) for p in payments)
    occupied = sum(1 for p in properties if p.status == "occupied")

    return jsonify(
        {
            "reports": list(REPORT_TYPES),
            "headline": {
                "properties": len(properties),
                "occupancy_rate": round(occupied / len(properties) * 100, 1) if properties else 0.0,
                "tenants": Tenant.query.count(),
                "active_leases": sum(1 for l in leases if l.status == "active"),
                "billed_to_date": billed,
                "collected_to_date": collected,
                "outstanding": billed - collected,
                "collection_rate": round(collected / billed * 100, 1) if billed else 0.0,
                "maintenance_total": len(requests),
                "maintenance_open": sum(1 for m in requests if m.is_open),
                "maintenance_cost": sum(float(m.cost or 0) for m in requests),
            },
        }
    )


@bp.get("")
@staff_required
def generate():
    report = request.args.get("report", "rent")
    try:
        filters = read_filters(request.args)
    except ValueError as exc:
        return jsonify({"message": str(exc)}), 400

    columns, rows, totals = build(report, filters)
    return jsonify(
        {
            "report": report,
            "label": next((r["label"] for r in REPORT_TYPES if r["value"] == report), report),
            "generated_at": date.today().isoformat(),
            "start_date": filters["start"].isoformat() if filters["start"] else None,
            "end_date": filters["end"].isoformat() if filters["end"] else None,
            "columns": columns,
            "rows": rows,
            "totals": totals,
            "money_columns": [c for c in columns if c in MONEY_COLUMNS],
        }
    )


def _title(report):
    return next((r["label"] for r in REPORT_TYPES if r["value"] == report), "Report")


def _range_label(filters):
    start = filters["start"].strftime("%d %b %Y") if filters["start"] else "Beginning"
    end = filters["end"].strftime("%d %b %Y") if filters["end"] else "Today"
    return f"{start} — {end}"


@bp.get("/export")
@staff_required
def export():
    """CSV, Excel or PDF export of any report."""
    report = request.args.get("report", "rent")
    fmt = request.args.get("format", "csv").lower()
    try:
        filters = read_filters(request.args)
    except ValueError as exc:
        return jsonify({"message": str(exc)}), 400

    columns, rows, totals = build(report, filters)
    stamp = date.today().isoformat()
    filename = f"propora-{report}-{stamp}"

    if fmt == "excel":
        return _excel(columns, rows, totals, report, filters, f"{filename}.xlsx")
    if fmt == "pdf":
        return _pdf(columns, rows, totals, report, filters, f"{filename}.pdf")

    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([_title(report)])
    writer.writerow([_range_label(filters)])
    writer.writerow([])
    writer.writerow(columns)
    writer.writerows(rows)
    writer.writerow([])
    for key, value in totals.items():
        writer.writerow([key, value])

    return Response(
        buffer.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}.csv"},
    )


def _excel(columns, rows, totals, report, filters, filename):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = report.title()[:31]

    heading = Font(bold=True, size=14)
    sheet["A1"] = _title(report)
    sheet["A1"].font = heading
    sheet["A2"] = _range_label(filters)
    sheet["A2"].font = Font(italic=True, color="666666")

    header_row = 4
    header_fill = PatternFill("solid", fgColor="2549C9")
    header_font = Font(bold=True, color="FFFFFF")
    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=header_row, column=index, value=column)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left")

    for r, row in enumerate(rows, start=header_row + 1):
        for c, value in enumerate(row, start=1):
            cell = sheet.cell(row=r, column=c, value=value)
            if columns[c - 1] in MONEY_COLUMNS and isinstance(value, (int, float)):
                cell.number_format = '₹#,##0'

    summary_row = header_row + len(rows) + 2
    sheet.cell(row=summary_row, column=1, value="Summary").font = Font(bold=True)
    for offset, (key, value) in enumerate(totals.items(), start=1):
        sheet.cell(row=summary_row + offset, column=1, value=key).font = Font(bold=True)
        sheet.cell(row=summary_row + offset, column=2, value=value)

    for index, column in enumerate(columns, start=1):
        width = max(len(str(column)), *(len(str(row[index - 1])) for row in rows)) if rows else len(column)
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width + 4, 12), 42)

    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return Response(
        stream.read(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


def _pdf(columns, rows, totals, report, filters, filename):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    stream = io.BytesIO()
    doc = SimpleDocTemplate(
        stream,
        pagesize=landscape(A4),
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title=_title(report),
    )
    styles = getSampleStyleSheet()
    cell_style = styles["BodyText"]
    cell_style.fontSize = 7.5
    cell_style.leading = 9.5

    story = [
        Paragraph(f"<b>PROPORA</b> — {_title(report)}", styles["Title"]),
        Paragraph(_range_label(filters), styles["Normal"]),
        Paragraph(f"Generated {date.today():%d %b %Y}", styles["Normal"]),
        Spacer(1, 8 * mm),
    ]

    def fmt(value, column):
        if column in MONEY_COLUMNS and isinstance(value, (int, float)):
            return f"₹{value:,.0f}"
        return str(value)

    table_data = [[Paragraph(f"<b>{c}</b>", cell_style) for c in columns]]
    # Keep the PDF readable — very long reports are truncated with a note.
    visible = rows[:400]
    for row in visible:
        table_data.append([Paragraph(fmt(v, columns[i]), cell_style) for i, v in enumerate(row)])

    table = Table(table_data, repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2549C9")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D6DBE4")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7FA")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(table)

    if len(rows) > len(visible):
        story.append(Spacer(1, 4 * mm))
        story.append(
            Paragraph(
                f"<i>Showing the first {len(visible)} of {len(rows)} rows. "
                f"Export to Excel or CSV for the complete data set.</i>",
                styles["Normal"],
            )
        )

    story.append(Spacer(1, 8 * mm))
    story.append(Paragraph("<b>Summary</b>", styles["Heading3"]))
    summary_table = Table(
        [[key, fmt(value, key)] for key, value in totals.items()],
        colWidths=[70 * mm, 50 * mm],
        hAlign="LEFT",
    )
    summary_table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D6DBE4")),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F5F7FA")),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.append(summary_table)

    doc.build(story)
    stream.seek(0)
    return Response(
        stream.read(),
        mimetype="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
